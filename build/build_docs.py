"""C391 LP + LG: convert the WSQ parent's built DOCX in place (design untouched).

Edits only: title, code, version record, WSQ layer (TSC, assessment, attendance,
funding), the 9:30am-5:30pm retiming, and slide refs remapped to the 106-slide deck.
"""
import re
from docx import Document

SRC = "source-wsq/courseware/{}-WSQ Technical Drawing with AutoCAD.docx"
OUT = "courseware/{}-AutoCAD Masterclass (C391).docx"
OLD_TITLE, NEW_TITLE = "WSQ Technical Drawing with AutoCAD", "AutoCAD Masterclass"
VERSION_ROW = ("1.0", "1 October 2026", None, "Dr Alfred Ang")

# 1-based source slides dropped from the deck (0-based 2,9,10,11,13,14,15,112,113,114)
DROPPED = {3, 10, 11, 12, 14, 15, 16, 113, 114, 115}
KEPT = [s for s in range(1, 117) if s not in DROPPED]
NEW_NO = {old: i + 1 for i, old in enumerate(KEPT)}


def remap(cell):
    def rng(m):
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        ks = [NEW_NO[s] for s in range(a, b + 1) if s in NEW_NO]
        return f"{ks[0]}–{ks[-1]}" if len(ks) > 1 and ks[0] != ks[-1] else str(ks[0])
    return re.sub(r"(\d+)(?:–(\d+))?", rng, cell)


def set_par(par, text):
    if not par.runs:
        par.add_run(text)
        return
    # cells may open with an EMPTY run; write into the first run that carries text
    runs = par.runs
    k = next((i for i, r in enumerate(runs) if r.text), 0)
    for i, r in enumerate(runs):
        r.text = text if i == k else ""


def drop(par):
    par._p.getparent().remove(par._p)


def set_cell(cell, text):
    pars = cell.paragraphs
    set_par(pars[0], text)
    for p in pars[1:]:
        drop(p)


def drop_row(table, row):
    table._tbl.remove(row._tr)


def version_record(doc, summary):
    t = doc.tables[0]
    for r in t.rows[2:]:
        drop_row(t, r)
    no, date, _, author = VERSION_ROW
    for c, v in zip(t.rows[1].cells, (no, date, summary, author)):
        set_cell(c, v)


def cover(doc):
    for p in doc.paragraphs[:14]:
        t = p.text
        if t == OLD_TITLE:
            set_par(p, NEW_TITLE)
        elif t.startswith("TGS Ref No"):
            set_par(p, "Course Code: C391")
        elif t.startswith("Version 13"):
            set_par(p, "Version 1.0")


def strip_ak(t):
    return re.sub(r"\s*\((?:A\d\s*[·,]\s*K\d)\)", "", t)


def schedule(table, plan):
    """plan: list of (match, (start, end, dur), new_text|None) in final order; rows not
    matched are deleted (tea breaks, assessment, attendance)."""
    rows = list(table.rows)[1:]
    keep = []
    for match, (start, end, dur), text in plan:
        row = next(r for r in rows if match in r.cells[2].text)
        rows.remove(row)
        c = row.cells
        set_cell(c[0], f"{start}–{end}")
        set_cell(c[1], f"{dur} min")
        if text:
            set_cell(c[2], text)
        if c[3].text.strip():
            set_cell(c[3], remap(c[3].text))
        keep.append(row)
    for r in rows:
        drop_row(table, r)


def times(start, durs):
    h, m = map(int, start.split(":"))
    t = h * 60 + m
    out = []
    for d in durs:
        out.append((f"{t // 60}:{t % 60:02d}", f"{(t + d) // 60}:{(t + d) % 60:02d}", d))
        t += d
    assert t == 17 * 60 + 30, f"day ends {t // 60}:{t % 60:02d}"
    return out


def lp():
    doc = Document(SRC.format("LP"))
    cover(doc)
    version_record(doc, "First release of the AutoCAD Masterclass: five topics and 16 hands-on "
                        "labs over two days, 9:30am – 5:30pm (7.5 instructional hours per day).")
    for p in list(doc.paragraphs):
        t = p.text
        if t.startswith("Assessment\t") or (p.style.name == "Heading 1" and t == "Assessment"):
            drop(p)
        elif p.style.name == "List Bullet" and re.search(
                r"Written Assessment|Practical Performance|Open Book|Final assessment|75% attendance", t):
            drop(p)
        elif "Design Documentation & Assessment" in t:
            set_par(p, t.replace("Design Documentation & Assessment", "Design Documentation")
                       .replace("Organizing Objects, ", "Organizing Objects & "))
        elif "(aligned to the Skills Framework TSC)" in t:
            set_par(p, t.replace(" (aligned to the Skills Framework TSC)", ""))
        elif t.startswith("Total training time"):
            set_par(p, "Day total: 7.5 instructional hours (9:30am – 5:30pm, 30-min lunch).")

    info = doc.tables[1]
    for r in list(info.rows):
        k = r.cells[0].text
        if k == "Course Title":
            set_cell(r.cells[1], NEW_TITLE)
        elif k == "WSQ Course Reference":
            set_cell(r.cells[0], "Course Code")
            set_cell(r.cells[1], "C391")
        elif k == "Skills Framework TSC":
            drop_row(info, r)
        elif k == "Duration":
            set_cell(r.cells[1], "2 days · 7.5 instructional hours per day (15 hours)")
        elif k == "Daily Timing":
            set_cell(r.cells[1], "9:30 am – 5:30 pm (30-min lunch)")

    d1 = times("9:30", [30, 45, 30, 60, 30, 120, 45, 105, 15])
    schedule(doc.tables[2], [
        ("Welcome", d1[0], "Welcome, course introduction, ground rules, trainer and learner "
                           "introductions, learning outcomes and course outline"),
        ("Topic 1", d1[1], None),
        ("Lab 1:", d1[2], None),
        ("Topic 2", d1[3], None),
        ("Lunch", d1[4], None),
        ("Lab 2:", d1[5], None),
        ("Topic 3", d1[6], None),
        ("Lab 6:", d1[7], None),
        ("Day 1 recap", d1[8], "Day 1 recap and Q&A"),
    ])
    # Day 2: the six topic/lab blocks rescaled 255 -> 420 min (x1.647, 5-min grain)
    d2 = times("9:30", [15, 65, 75, 80, 30, 65, 75, 60, 15])
    schedule(doc.tables[3], [
        ("Day 1 recap", d2[0], "Day 1 recap"),
        ("Topic 4", d2[1], None),
        ("Lab 9:", d2[2], None),
        ("Lab 11:", d2[3], None),
        ("Lunch", d2[4], None),
        ("Topic 5", d2[5], None),
        ("Lab 13:", d2[6], None),
        ("Lab 15:", d2[7], None),
        ("Course recap", d2[8], "Course recap, Q&A and course feedback"),
    ])

    # Lab Reference: drop the TSC Ability · Knowledge column
    lab = doc.tables[4]
    for r in lab.rows:
        tc = r._tr.tc_lst[1]
        w_from = tc.tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcW")
        tc0 = r._tr.tc_lst[0].tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcW")
        if w_from is not None and tc0 is not None:
            q = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"
            tc0.set(q, str(int(tc0.get(q)) + int(w_from.get(q))))
        r._tr.remove(tc)
    grid = lab._tbl.tblGrid
    cols = grid.gridCol_lst
    cols[0].w = cols[0].w + cols[1].w
    grid.remove(cols[1])
    set_cell(lab.rows[0].cells[0], "Topic")
    doc.core_properties.title = NEW_TITLE
    doc.save(OUT.format("LP"))


LG_SUBS = [
    ("the WSQ course WSQ Technical Drawing with AutoCAD (TGS-2021005538)",
     "the course AutoCAD Masterclass (C391)"),
    (", aligned to the Skills Framework TSC Technical Drawing (DSN-TDR-3005-1.11)", ""),
    (" The final assessment is open book: you may refer to the slides, this Learner Guide "
     "and other approved materials.", ""),
    ("These conventions (K1/K3) apply across every lab and are assessed in the written assessment.",
     "These conventions apply across every lab."),
    ("Preparing for the Assessment", "Consolidating Your Skills"),
    ("Review the Key Concepts bullets of each topic — the written assessment (SAQ) draws on them.",
     "Review the Key Concepts bullets of each topic to consolidate the theory behind each lab."),
    ("The practical performance asks you to produce a dimensioned technical drawing end-to-end: "
     "set up, draw, modify, organise with layers, dimension, and plot — Labs 2, 7, 9, 14 and 16 "
     "cover that chain.",
     "Practise producing a dimensioned technical drawing end-to-end: set up, draw, modify, "
     "organise with layers, dimension, and plot — Labs 2, 7, 9, 14 and 16 cover that chain."),
    ("Course materials, assessment submission and your certificate are on the LMS:",
     "Course materials and your certificate of completion are on the LMS:"),
]


def lg_text(t):
    for a, b in LG_SUBS:
        t = t.replace(a, b)
    t = strip_ak(t)
    return re.sub(r" \(A\d, K\d\)\.", ".", t)


def lg():
    doc = Document(SRC.format("LG"))
    cover(doc)
    version_record(doc, "First release of the AutoCAD Masterclass learner guide: 16 "
                        "step-by-step labs, a drawing-reading reference and a glossary.")
    for p in doc.paragraphs:
        new = lg_text(p.text)
        if new != p.text:
            set_par(p, new)
    doc.core_properties.title = NEW_TITLE
    doc.save(OUT.format("LG"))


def lg_md():
    src = open(f"source-wsq/LG-{OLD_TITLE}.md", encoding="utf-8").read()
    src = src.replace(f"# {OLD_TITLE} — Learner Guide", f"# {NEW_TITLE} — Learner Guide")
    src = src.replace("**WSQ Course Code:** TGS-2021005538", "**Course Code:** C391")
    src = src.replace("**Version v13 · 21 August 2026**", "**Version v1.0 · 1 October 2026**")
    src = src.replace("(#preparing-for-the-assessment)", "(#consolidating-your-skills)")
    src = re.sub(r"--a(\d)--k\d\)", ")", src)  # TOC anchors for the stripped (A# · K#)
    out = "\n".join(lg_text(line) for line in src.split("\n"))
    open(f"LG-{NEW_TITLE} (C391).md", "w", encoding="utf-8").write(out)


if __name__ == "__main__":
    lp()
    lg()
    lg_md()
    print("docs ok")
