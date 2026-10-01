"""C391 deck: residual WSQ-layer edits the converter does not cover."""
import re
from pptx import Presentation
F = "courseware/AutoCAD Masterclass (C391)-v1.0.pptx"
p = Presentation(F)

def set_para(par, text):
    runs = par.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""

def drop_para(par):
    par._p.getparent().remove(par._p)

def shape(slide, sid):
    return next(sh for sh in slide.shapes if sh.shape_id == sid)

def edit(slide, sid, fn):
    tf = shape(slide, sid).text_frame
    for par in list(tf.paragraphs):
        new = fn(par.text)
        if new is None:
            drop_para(par)
        elif new != par.text:
            set_para(par, new)

S = p.slides
# cover: no TSC line on a non-WSQ course
edit(S[0], 12, lambda t: None if t.startswith("Skills Framework") else t)
# trainer profile
edit(S[3], 21, lambda t: t.replace("WSQ courses on", "Courses on"))
# schedule slide
edit(S[7], 5, lambda t: t.replace("8 hours/day", "7.5 hours/day"))
edit(S[7], 11, lambda t: t.replace("Digital attendance (AM) · introductions · outcomes",
                                   "Welcome · introductions · course overview"))
def day2(t):
    if t.startswith("–  Final Assessment"): return None
    t = t.replace("Organizing Objects, Design Documentation & Assessment",
                  "Organizing Objects & Design Documentation")
    t = t.replace("Course feedback & TRAQOM survey", "Course summary & feedback")
    t = t.replace("9:30am–6:30pm · 1-hour lunch", "9:30am–5:30pm · 30-min lunch")
    return t
edit(S[7], 12, day2)
# course outline: strip TSC A#/K# codes
for sid in (11, 12):
    edit(S[8], sid, lambda t: re.sub(r" \(A\d, K\d\)", "", t))
edit(S[10], 4, lambda t: t.replace("K1 · K3 — ", ""))
for i, n in ((13, 1), (28, 2), (46, 3), (63, 4), (81, 5)):
    edit(S[i], 4, lambda t, n=n: f"TOPIC 0{n} · KEY CONCEPTS" if t.startswith("SKILLS FRAMEWORK") else t)
edit(S[101], 4, lambda t: "COURSE RECAP" if t == "LEARNING OUTCOMES" else t)
# TRAQOM slide -> Course Feedback (rewrite in place)
edit(S[104], 5, lambda t: "Course Feedback")
fb = ["•  Please complete the course feedback form on the LMS — it helps us improve the course.",
      "•  Your certificate of completion is issued on the LMS after the course.",
      "•  Portal: https://lms-tms.tertiaryinfotech.com"]
pars = shape(S[104], 7).text_frame.paragraphs
for par, txt in zip(pars, fb):
    set_para(par, txt)
p.save(F)
print("ok")

# pass 2: Download Course Material — no assessment in a non-WSQ course
p = Presentation(F)
for sh in p.slides[6].shapes:
    if sh.has_text_frame:
        for par in sh.text_frame.paragraphs:
            if "final assessment is open book" in par.text:
                set_para(par, "•  Keep them handy — refer to the slides and Learner Guide during the labs.")
p.save(F)
print("pass2 ok")
